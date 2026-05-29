"""DOCX document export functionality."""

from pathlib import Path
from datetime import datetime
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from app.brd import BRD
from app.utils.logger import get_logger

logger = get_logger(__name__)


class DOCXExporter:
    """Export BRD to Word (.docx) format."""

    @staticmethod
    def export(brd: BRD, output_path: Path) -> Path:
        """Export BRD to DOCX."""
        logger.info(f"Exporting BRD to {output_path}")
        
        doc = Document()
        
        # Add title
        title = doc.add_heading(brd.title, level=1)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # Add metadata
        doc.add_heading("Document Information", level=2)
        metadata_table = doc.add_table(rows=3, cols=2)
        metadata_table.rows[0].cells[0].text = "Document ID"
        metadata_table.rows[0].cells[1].text = brd.document_id
        metadata_table.rows[1].cells[0].text = "Version"
        metadata_table.rows[1].cells[1].text = brd.version
        metadata_table.rows[2].cells[0].text = "Generated"
        metadata_table.rows[2].cells[1].text = datetime.now().isoformat()
        
        # Executive Summary
        if brd.executive_summary:
            doc.add_heading("Executive Summary", level=2)
            doc.add_paragraph(brd.executive_summary)
        
        # Business Objectives
        if brd.business_objectives:
            doc.add_heading("Business Objectives", level=2)
            for obj in brd.business_objectives:
                doc.add_paragraph(obj, style="List Bullet")
        
        # Scope
        doc.add_heading("Scope", level=2)
        if brd.scope_in:
            doc.add_heading("In Scope", level=3)
            for item in brd.scope_in:
                doc.add_paragraph(item, style="List Bullet")
        
        if brd.scope_out:
            doc.add_heading("Out of Scope", level=3)
            for item in brd.scope_out:
                doc.add_paragraph(item, style="List Bullet")
        
        # Functional Requirements
        if brd.functional_requirements:
            doc.add_heading("Functional Requirements", level=2)
            fr_table = doc.add_table(rows=1, cols=4)
            fr_table.style = "Light Grid Accent 1"
            hdr_cells = fr_table.rows[0].cells
            hdr_cells[0].text = "ID"
            hdr_cells[1].text = "Requirement"
            hdr_cells[2].text = "Priority"
            hdr_cells[3].text = "Owner"
            
            for fr in brd.functional_requirements:
                row_cells = fr_table.add_row().cells
                row_cells[0].text = fr.id
                row_cells[1].text = fr.requirement
                row_cells[2].text = fr.priority.value if fr.priority else "TBD"
                row_cells[3].text = fr.owner or "TBD"
        
        # Non-Functional Requirements
        if brd.nonfunctional_requirements:
            doc.add_heading("Non-Functional Requirements", level=2)
            nfr_table = doc.add_table(rows=1, cols=3)
            nfr_table.style = "Light Grid Accent 1"
            hdr_cells = nfr_table.rows[0].cells
            hdr_cells[0].text = "Category"
            hdr_cells[1].text = "Requirement"
            hdr_cells[2].text = "Priority"
            
            for nfr in brd.nonfunctional_requirements:
                row_cells = nfr_table.add_row().cells
                row_cells[0].text = nfr.category or "General"
                row_cells[1].text = nfr.requirement
                row_cells[2].text = nfr.priority.value if nfr.priority else "TBD"
        
        # Business Rules
        if brd.business_rules:
            doc.add_heading("Business Rules", level=2)
            for rule in brd.business_rules:
                doc.add_paragraph(rule, style="List Bullet")
        
        # Risks
        if brd.risks:
            doc.add_heading("Risks & Open Items", level=2)
            for risk in brd.risks:
                doc.add_paragraph(risk, style="List Bullet")
        
        # Acceptance Criteria
        if brd.acceptance_criteria:
            doc.add_heading("Acceptance Criteria", level=2)
            for criterion in brd.acceptance_criteria:
                doc.add_paragraph(criterion, style="List Bullet")
        
        # Save
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        logger.info(f"BRD exported to {output_path}")
        return output_path
